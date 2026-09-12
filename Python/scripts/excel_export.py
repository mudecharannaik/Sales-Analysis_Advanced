"""Excel export utilities.

Creates a comprehensive Excel workbook with multiple formatted sheets
summarising the analysis results.
"""

from __future__ import annotations

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import PROJECT_ROOT, get_logger

logger = get_logger(__name__)

EXCEL_DIR = PROJECT_ROOT / "Excel"
EXCEL_DIR.mkdir(parents=True, exist_ok=True)


def _auto_width(ws) -> None:
    """Auto-fit column widths based on cell contents.

    Args:
        ws: openpyxl worksheet.
    """
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            cell.alignment = Alignment(wrap_text=True)
            val = str(cell.value) if cell.value is not None else ""
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _format_header(ws, header_row: int = 1) -> None:
    """Apply bold, coloured header formatting.

    Args:
        ws: openpyxl worksheet.
        header_row: Row number containing headers.
    """
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _write_dataframe(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    """Write a DataFrame to a worksheet with headers.

    Args:
        ws: openpyxl worksheet.
        df: DataFrame to write.
        start_row: Starting row for the data.
    """
    headers = list(df.columns)
    for c_idx, header in enumerate(headers, start=1):
        ws.cell(row=start_row, column=c_idx, value=str(header))

    for r_idx, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for c_idx, value in enumerate(row, start=1):
            if pd.isna(value):
                ws.cell(row=r_idx, column=c_idx, value="")
            else:
                ws.cell(row=r_idx, column=c_idx, value=value)


def create_excel_report(
    quality_profile: pd.DataFrame,
    eda_summary: pd.DataFrame,
    stat_summary: pd.DataFrame,
    cleaned_df: pd.DataFrame,
    output_name: str = "SalesAnalysis_Report.xlsx",
) -> str:
    """Create a comprehensive Excel workbook.

    Sheets:
        - Data_Quality_Report: column profiling output
        - EDA_Summary: exploratory data analysis summary
        - Statistical_Summary: statistical test results
        - Cleaned_Sample: first 1000 rows of cleaned data

    Args:
        quality_profile: DataFrame from data quality profiling.
        eda_summary: DataFrame with EDA summary statistics.
        stat_summary: DataFrame with statistical test results.
        cleaned_df: Cleaned DataFrame (sample will be used).
        output_name: Output filename.

    Returns:
        Absolute path of the saved workbook.
    """
    logger.info("Creating Excel report: %s", output_name)

    wb = Workbook()

    # Data Quality Report
    ws_quality = wb.active
    ws_quality.title = "Data_Quality_Report"
    _write_dataframe(ws_quality, quality_profile.reset_index())
    _format_header(ws_quality)
    _auto_width(ws_quality)

    # EDA Summary
    ws_eda = wb.create_sheet("EDA_Summary")
    _write_dataframe(ws_eda, eda_summary.reset_index() if eda_summary.index.name else eda_summary)
    _format_header(ws_eda)
    _auto_width(ws_eda)

    # Statistical Summary
    ws_stat = wb.create_sheet("Statistical_Summary")
    _write_dataframe(ws_stat, stat_summary)
    _format_header(ws_stat)
    _auto_width(ws_stat)

    # Cleaned Sample
    ws_clean = wb.create_sheet("Cleaned_Sample")
    sample = cleaned_df.head(1000)
    _write_dataframe(ws_clean, sample)
    _format_header(ws_clean)
    _auto_width(ws_clean)

    output_path = EXCEL_DIR / output_name
    wb.save(output_path)
    logger.info("Excel report saved to %s", output_path)
    return str(output_path)
